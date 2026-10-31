"""
Tests for Kora, Kura's AI assistant.

    python manage.py test kura.test_agent

Covers the deterministic engine (charts, tables, cards, maps, shapes,
cleaning, privacy, dashboards, questionnaires), the privacy promise that no
raw answer reaches the language model, the built-in planner, every HTTP
endpoint with permissions, the file exports and the page injection rules.
"""

import io
import json
import random
from datetime import timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone

from . import agent_ai, agent_engine as eng
from .models import FormVersion, Submission, Survey

User = get_user_model()

REGIONS = ["Banjul", "Kanifing", "West Coast", "North Bank", "Lower River", "Central River", "Upper River"]
SCHEMA = {
    "settings": {"anonymous": True},
    "questions": [
        {"name": "intro", "type": "note", "label": "Welcome"},
        {"name": "consent", "type": "select_one", "label": "Do you agree?", "required": True,
         "choices": [{"value": "yes", "label": "Yes"}, {"value": "no", "label": "No"}]},
        {"name": "full_name", "type": "text", "label": "Respondent's full name"},
        {"name": "phone", "type": "text", "label": "Phone number"},
        {"name": "age", "type": "integer", "label": "How old are you?", "required": True},
        {"name": "sex", "type": "select_one", "label": "Sex",
         "choices": [{"value": "f", "label": "Female"}, {"value": "m", "label": "Male"}]},
        {"name": "region", "type": "select_one", "label": "Region",
         "choices": [{"value": r.lower().replace(" ", "_"), "label": r} for r in REGIONS]},
        {"name": "water", "type": "select_one", "label": "Main water source",
         "choices": [{"value": "tap", "label": "Tap"}, {"value": "well", "label": "Well"},
                     {"value": "borehole", "label": "Borehole"}]},
        {"name": "assets", "type": "select_multiple", "label": "Assets owned",
         "choices": [{"value": "radio", "label": "Radio"}, {"value": "tv", "label": "TV"},
                     {"value": "phone_set", "label": "Mobile phone"}]},
        {"name": "satisfaction", "type": "likert", "label": "I am satisfied",
         "choices": [{"value": str(i), "label": l} for i, l in
                     enumerate(["Strongly disagree", "Disagree", "Neutral", "Agree", "Strongly agree"], 1)]},
        {"name": "income", "type": "decimal", "label": "Monthly income"},
        {"name": "village", "type": "text", "label": "Village name"},
        {"name": "hh_gps", "type": "gps", "label": "Household location"},
        {"name": "comments", "type": "long_text", "label": "Any comments?"},
        {"name": "members", "type": "repeat", "label": "Household members",
         "children": [{"name": "m_age", "type": "integer", "label": "Member age"}]},
        {"name": "bmi", "type": "calculate", "label": "BMI", "calc": "1"},
    ],
}
SECRET_NAMES = ["Fatou Ceesay", "Lamin Jallow", "Awa Touray", "Ebrima Sowe"]


def make_survey(owner, n=60, seed=7):
    rnd = random.Random(seed)
    s = Survey.objects.create(owner=owner, title="Household test survey", draft_schema=SCHEMA, state="collecting")
    fv = FormVersion.objects.create(survey=s, version=1, schema=SCHEMA, is_current=True, published_by=owner)
    now = timezone.now()
    for i in range(n):
        ans = {
            "consent": "yes",
            "full_name": SECRET_NAMES[i % 4] + ("  " if i % 9 == 0 else ""),
            "phone": f"+220 7{i:06d}",
            "age": 18 + (i * 7) % 60,
            "sex": "f" if i % 3 else "m",
            "region": REGIONS[i % 7].lower().replace(" ", "_"),
            "water": ["tap", "well", "borehole"][i % 3],
            "assets": ["radio", "tv"] if i % 2 else ["phone_set"],
            "satisfaction": str(1 + i % 5),
            "income": 2000 + rnd.randint(0, 3000) if i != 5 else 9_999_999,
            "village": " Brikama" if i % 4 == 0 else "Brikama",
            "hh_gps": f"{13.2 + i * 0.01:.5f} {-16.6 + i * 0.01:.5f}",
            "comments": "My neighbour Fatou Ceesay helped" if i == 3 else "",
            "members": [{"m_age": 30}, {"m_age": 4}],
        }
        sub = Submission.objects.create(survey=s, form_version=fv, answers=ans, calculations={"bmi": 20 + i % 5},
                                        status="complete", source="app" if i % 2 else "web",
                                        gps_lat=13.4, gps_lng=-16.6, duration_ms=300000 + i * 1000)
        Submission.objects.filter(pk=sub.pk).update(received_at=now - timedelta(hours=i * 5))
    # an exact duplicate submission
    dup = Submission.objects.filter(survey=s).first()
    Submission.objects.create(survey=s, form_version=fv, answers=dup.answers, calculations=dup.calculations,
                              status="complete", source="web")
    return s


class AgentEngineTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("owner", password="pw", first_name="Baboucarr")
        cls.survey = make_survey(cls.owner)

    def ctx(self, page="data"):
        from .agent_views import build_context
        return build_context(self.owner, self.survey.code, page)

    def test_catalog_uses_frame_column_names(self):
        ctx = self.ctx()
        names = {f["name"] for f in ctx.catalog}
        self.assertIn("calc__bmi", names)
        self.assertNotIn("intro", names)
        self.assertIn("calc__bmi", ctx.frame.columns)

    def test_pii_classification(self):
        ctx = self.ctx()
        det = {f["name"]: eng.detect_pii(f, ctx.frame.get(f["name"])) for f in ctx.catalog}
        self.assertEqual(det["full_name"]["level"], "direct")
        self.assertEqual(det["full_name"]["action"], "drop")
        self.assertEqual(det["phone"]["level"], "direct")
        self.assertEqual(det["hh_gps"]["action"], "drop")
        self.assertEqual(det["age"], {"level": "quasi", "reason": "age", "action": "band10"})
        self.assertEqual(det["comments"]["action"], "drop")
        self.assertEqual(det["village"]["level"], "quasi")      # "village name" is a place, not a person
        self.assertEqual(det["water"]["level"], "none")

    def test_prompt_never_contains_raw_answers(self):
        ctx = self.ctx()
        prompt = agent_ai.build_user_prompt(ctx, "tell me about my data", [])
        for name in SECRET_NAMES:
            self.assertNotIn(name, prompt)
        self.assertNotIn("+220 7000001", prompt)
        self.assertNotIn("Brikama", prompt)
        self.assertNotIn("13.21", prompt)
        self.assertIn("Female", prompt)                    # aggregates are allowed
        self.assertIn("⚠ direct", prompt)

    def test_chart_bar_uses_labels_and_percentages(self):
        b = eng.run_tool(self.ctx(), {"tool": "chart", "kind": "bar", "x": "region"})
        self.assertEqual(b["type"], "chart", b)
        self.assertTrue(set(b["labels"]) <= set(REGIONS))
        self.assertAlmostEqual(sum(b["percent"]), 100, delta=1)
        self.assertIn("Most common", b["caption"])

    def test_chart_likert_keeps_form_order(self):
        b = eng.run_tool(self.ctx(), {"tool": "chart", "kind": "bar", "x": "satisfaction"})
        self.assertEqual(b["labels"][0], "Strongly disagree")
        self.assertEqual(b["labels"][-1], "Strongly agree")

    def test_chart_multi_select_is_exploded(self):
        b = eng.run_tool(self.ctx(), {"tool": "chart", "kind": "hbar", "x": "assets"})
        self.assertEqual(set(b["labels"]), {"Radio", "TV", "Mobile phone"})
        self.assertIn("more than one answer", b["caption"])

    def test_chart_variants(self):
        ctx = self.ctx()
        g = eng.run_tool(ctx, {"tool": "chart", "kind": "stacked", "x": "region", "group_by": "sex"})
        self.assertEqual(len(g["series"]), 2)
        m = eng.run_tool(ctx, {"tool": "chart", "kind": "bar", "x": "sex", "y": "age", "agg": "mean"})
        self.assertEqual(m["type"], "chart")
        hst = eng.run_tool(ctx, {"tool": "chart", "kind": "histogram", "x": "age"})
        self.assertIn("Mean", hst["caption"])
        sc = eng.run_tool(ctx, {"tool": "chart", "kind": "scatter", "x": "age", "y": "calc__bmi"})
        self.assertIn("Correlation", sc["caption"])
        ts = eng.run_tool(ctx, {"tool": "chart", "kind": "timeseries", "freq": "D"})
        self.assertEqual(ts["chart"], "line")
        self.assertTrue(ts["labels"])

    def test_filters_accept_choice_labels(self):
        b = eng.run_tool(self.ctx(), {"tool": "cards", "cards": [
            {"label": "Women", "agg": "count", "filters": [{"field": "sex", "operator": "eq", "value": "Female"}]},
            {"label": "Share of women", "format": "percent", "filters": [{"field": "sex", "operator": "eq", "value": "Female"}]},
            {"label": "Has a radio", "agg": "count", "filters": [{"field": "assets", "operator": "eq", "value": "Radio"}]},
        ]})
        women, share, radio = b["cards"]
        df = self.ctx().frame
        self.assertEqual(women["value"], int((df["sex"] == "f").sum()))
        self.assertAlmostEqual(share["value"], round((df["sex"] == "f").mean() * 100, 1))
        self.assertEqual(radio["value"], int(df["assets"].apply(lambda v: "radio" in v).sum()))

    def test_tables(self):
        ctx = self.ctx()
        f = eng.run_tool(ctx, {"tool": "table", "mode": "frequency", "x": "water"})
        self.assertEqual(f["columns"], ["Main water source", "Responses", "Percent"])
        c = eng.run_tool(ctx, {"tool": "table", "mode": "crosstab", "x": "region", "group_by": "sex", "percent": "row"})
        self.assertIn("Female", c["columns"])
        g = eng.run_tool(ctx, {"tool": "table", "mode": "group", "group_by": "water", "y": "age"})
        self.assertIn("Mean", g["columns"])
        r = eng.run_tool(ctx, {"tool": "table", "mode": "rows", "columns": ["sex", "assets"], "limit": 5})
        self.assertEqual(len(r["rows"]), 5)
        self.assertTrue(r["row_level"])
        self.assertIn(r["rows"][0]["Sex"], ("Female", "Male"))
        sm = eng.run_tool(ctx, {"tool": "table", "mode": "summary"})
        self.assertIn("Median", sm["columns"])

    def test_map_and_shapes(self):
        ctx = self.ctx()
        m = eng.run_tool(ctx, {"tool": "map", "color_by": "water"})
        self.assertEqual(m["type"], "map", m)
        self.assertEqual(set(m["legend"]), {"Tap", "Well", "Borehole"})
        r = eng.run_tool(ctx, {"tool": "shape", "kind": "ring", "x": "sex"})
        self.assertEqual(len(r["items"]), 2)
        fn = eng.run_tool(ctx, {"tool": "shape", "kind": "funnel", "items": [
            {"label": "Everyone"}, {"label": "Women", "filters": [{"field": "sex", "operator": "eq", "value": "f"}]},
            {"label": "Women with a TV", "filters": [{"field": "sex", "operator": "eq", "value": "f"},
                                                     {"field": "assets", "operator": "eq", "value": "tv"}]}]})
        vals = [i["value"] for i in fn["items"]]
        self.assertEqual(vals, sorted(vals, reverse=True))

    def test_profile_finds_problems(self):
        b = eng.run_tool(self.ctx(), {"tool": "profile"})
        text = " ".join(i["text"] for i in b["issues"])
        self.assertGreaterEqual(b["duplicates"], 1)
        self.assertIn("extreme", text)
        self.assertIn("stray spaces", text)
        self.assertIn("personal data", text)

    def test_auto_clean(self):
        ctx = self.ctx()
        b = eng.run_tool(ctx, {"tool": "clean", "auto": True})
        ops = [s["operation"] for s in b["steps"]]
        self.assertIn("trim_text", ops)
        self.assertIn("deduplicate", ops)
        self.assertTrue(all(s["status"] == "ok" for s in b["steps"]), b["steps"])
        self.assertLess(b["after"]["rows"], b["before"]["rows"])
        clean, _ = eng.run_clean(ctx, b["spec"]["steps"])
        self.assertFalse(clean["village"].astype(str).str.startswith(" ").any())
        # originals untouched
        self.assertTrue(Submission.objects.filter(answers__village=" Brikama").exists())

    def test_clean_rejects_unknown_operations(self):
        b = eng.run_tool(self.ctx(), {"tool": "clean", "steps": [{"operation": "rm_rf", "config": {}}]})
        # an unknown operation is refused, never swapped for something else
        self.assertEqual(b["type"], "error")
        self.assertIn("cleaning steps", b["text"])

    def test_privacy_copy(self):
        ctx = self.ctx()
        out, plan, k = eng.run_privacy(ctx, {})
        cols = set(out.columns)
        for gone in ("full_name", "hh_gps", "comments", "_uuid", "_gps_lat", "members"):
            self.assertNotIn(gone, cols)
        self.assertIn("phone", cols)
        self.assertTrue(out["phone"].dropna().str.startswith("id_").all())
        self.assertTrue(out["age"].dropna().str.contains("–").all())
        flat = out.astype(str).to_string()
        for name in SECRET_NAMES:
            self.assertNotIn(name, flat)
        self.assertIn("row_id", cols)
        self.assertGreater(k["at_risk"], 0)
        out2, _p, k2 = eng.run_privacy(ctx, {"suppress": True, "k": 3})
        self.assertEqual(k2["suppressed"], k2["at_risk"])
        self.assertLess(len(out2), len(out))
        out3, _p, _k = eng.run_privacy(ctx, {"rules": {"hh_gps": "round2", "phone": "drop"}})
        self.assertNotIn("phone", out3.columns)
        self.assertTrue(out3["hh_gps"].str.fullmatch(r"\d+\.\d{1,2} -\d+\.\d{1,2}").all(), out3["hh_gps"].head())

    def test_pseudonyms_are_stable_and_one_way(self):
        a = eng._pseudo("ABC", "phone", "+2207000001")
        self.assertEqual(a, eng._pseudo("ABC", "phone", "+2207000001"))
        self.assertNotEqual(a, eng._pseudo("XYZ", "phone", "+2207000001"))
        self.assertNotIn("7000001", a)

    def test_auto_dashboard(self):
        b = eng.run_tool(self.ctx(), {"tool": "dashboard", "auto": True, "theme": "ocean"})
        self.assertEqual(b["type"], "dashboard")
        self.assertEqual(b["theme"], "ocean")
        kinds = {it["block"]["type"] for it in b["items"]}
        self.assertTrue({"cards", "chart", "map"} <= kinds, kinds)
        self.assertFalse([it for it in b["items"] if it["block"]["type"] == "error"],
                         [it["block"] for it in b["items"] if it["block"]["type"] == "error"])
        for it in b["items"]:
            self.assertLessEqual(it["x"] + it["w"], 12)

    def test_questionnaire_sanitizer(self):
        qs = eng.sanitize_questions([
            {"type": "radio", "label": "Do you agree?", "name": "consent", "choices": ["Yes", "No"]},
            {"type": "number", "label": "Age", "name": "consent"},                         # duplicate name
            {"type": "weird", "label": "Something"},                                      # unknown type
            {"type": "text", "label": "Why not?", "relevant": {"op": "and", "rules": [
                {"q": "consent", "cmp": "eq", "value": "no"}, {"q": "later_q", "cmp": "eq", "value": 1}]}},
            {"type": "select_one", "label": "Later", "name": "later_q", "choices": []},
            {"type": "repeat", "label": "Members", "children": [{"type": "integer", "label": "Member age"},
                                                                {"type": "repeat", "label": "nested"}]},
        ])
        self.assertEqual(qs[0]["type"], "select_one")
        self.assertEqual([c["label"] for c in qs[0]["choices"]], ["Yes", "No"])
        self.assertEqual(qs[1]["type"], "integer")
        self.assertNotEqual(qs[1]["name"], "consent")
        self.assertEqual(qs[2]["type"], "text")
        self.assertEqual(qs[3]["relevant"]["rules"], [{"q": "consent", "cmp": "eq", "value": "no"}])
        self.assertEqual(len(qs[4]["choices"]), 2)                                         # filled in
        self.assertEqual(qs[5]["children"][1]["type"], "text")                            # no nested repeats
        names = [q["name"] for q in qs] + [c["name"] for c in qs[5]["children"]]
        self.assertEqual(len(names), len(set(names)))

    def test_review_form(self):
        sch = json.loads(json.dumps(SCHEMA))
        sch["questions"].append({"name": "pick", "type": "select_one", "label": "Pick one",
                                 "choices": [{"value": "option_1", "label": "Option 1"}]})
        sch["questions"].append({"name": "late", "type": "text", "label": "Late",
                                 "relevant": {"op": "and", "rules": [{"q": "nowhere", "cmp": "eq", "value": 1}]}})
        ctx = self.ctx("builder")
        ctx.draft_schema = sch
        b = eng.run_tool(ctx, {"tool": "review_form"})
        text = " ".join(i["text"] for i in b["issues"])
        self.assertIn("placeholder", text)
        self.assertIn("fewer than two", text)
        self.assertIn("doesn't come earlier", text)

    def test_review_groups_repeated_advice(self):
        b = eng.run_tool(self.ctx(), {"tool": "review_form"})
        ranges = [i for i in b["issues"] if "minimum or maximum" in i["text"]]
        self.assertEqual(len(ranges), 1)                      # one grouped line, not one per question
        self.assertFalse([i for i in b["issues"] if "Do you agree" in i["text"] and "Other" in i["text"]])

    def test_dashboard_rows_have_no_holes(self):
        tiles = eng.layout_tiles([{"kind": "kpi"}, {"kind": "kpi"}, {"kind": "kpi"},
                                  {"kind": "timeseries"}, {"kind": "pie"}, {"kind": "hbar"}])
        rows = {}
        for t in tiles:
            rows.setdefault(t["y"], []).append(t)
        for row in rows.values():
            self.assertEqual(sum(t["w"] for t in row), 12)
            self.assertEqual(len({t["h"] for t in row}), 1)

    def test_whole_number_histogram_bands(self):
        b = eng.run_tool(self.ctx(), {"tool": "chart", "kind": "histogram", "x": "age"})
        self.assertRegex(b["labels"][0], r"^\d+–\d+$")

    def test_clean_blank_share_ignores_marker_columns(self):
        b = eng.run_tool(self.ctx(), {"tool": "clean", "auto": True})
        self.assertLessEqual(b["after"]["missing_pct"], b["before"]["missing_pct"] + 0.5)

    def test_tool_errors_are_friendly(self):
        b = eng.run_tool(self.ctx(), {"tool": "chart", "x": "no_such_column"})
        self.assertEqual(b["type"], "error")
        self.assertIn("no_such_column", b["text"])
        self.assertEqual(eng.run_tool(self.ctx(), {"tool": "launch_rockets"})["type"], "error")


class BuiltinPlannerTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("owner2", password="pw")
        cls.survey = make_survey(cls.owner, n=20)

    def plan(self, msg, code=True, page="data"):
        from .agent_views import build_context
        ctx = build_context(self.owner, self.survey.code if code else None, page)
        return agent_ai.normalise_plan(agent_ai.builtin_plan(ctx, msg)), ctx

    def tools(self, msg, **kw):
        p, _ctx = self.plan(msg, **kw)
        return [a["tool"] for a in p["actions"]]

    def test_intents(self):
        self.assertEqual(self.tools("make me a dashboard"), ["dashboard"])
        self.assertEqual(self.tools("clean my data"), ["clean"])
        self.assertEqual(self.tools("make a privacy safe copy to share"), ["privacy"])
        self.assertEqual(self.tools("any data problems?"), ["profile"])
        self.assertIn("overview", self.tools("how is collection going today?"))
        self.assertEqual(self.tools("map the responses by water source"), ["map"])
        self.assertEqual(self.tools("draft a questionnaire about clinic waiting times"), ["questionnaire"])
        self.assertEqual(self.tools("key numbers as cards"), ["cards"])
        self.assertEqual(self.tools("review my form", page="builder"), ["review_form"])

    def test_chart_picks_the_right_column(self):
        p, _ = self.plan("pie chart of water source")
        a = p["actions"][0]
        self.assertEqual((a["tool"], a["kind"], a["x"]), ("chart", "pie", "water"))
        p, _ = self.plan("chart region by sex")
        self.assertEqual(p["actions"][0].get("group_by"), "sex")
        p, _ = self.plan("table of region by sex in percent")
        self.assertEqual(p["actions"][0]["mode"], "crosstab")

    def test_no_survey(self):
        p, _ = self.plan("chart something", code=False, page="list")
        self.assertEqual(p["actions"], [])
        self.assertEqual(self.tools("show my surveys", code=False, page="list"), ["surveys"])

    def test_template_skips_questions_the_form_already_has(self):
        p, _ = self.plan("draft a questionnaire about household water", page="builder")
        names = [q["name"] for q in p["actions"][0]["questions"]]
        for already in ("consent", "age", "sex", "region"):
            self.assertNotIn(already, names)
        # every skip rule still points at a real question: the form's own consent, or one in the template
        known = set(names) | {q["name"] for q in SCHEMA["questions"]}
        rules = [r for q in p["actions"][0]["questions"] for r in (q.get("relevant") or {}).get("rules", [])]
        self.assertTrue(rules)
        self.assertTrue(all(r["q"] in known for r in rules))
        self.assertIn("consent", {r["q"] for r in rules})

    def test_templates(self):
        for topic, title in [("household", "Household survey"), ("school", "School and learning survey"),
                             ("water and sanitation", "WASH survey"), ("farming", "Farm and harvest survey")]:
            q = agent_ai.template_questionnaire(f"draft a survey about {topic}")
            self.assertEqual(q["title"], title)
            qs = eng.sanitize_questions(q["questions"])
            self.assertEqual(qs[1]["name"], "consent")
            self.assertTrue(all(x["relevant"] for x in qs[2:]))     # gated on consent


class AgentViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("own", password="pw", first_name="Baboucarr")
        cls.outsider = User.objects.create_user("out", password="pw")
        cls.viewer = User.objects.create_user("viewer", password="pw")
        cls.survey = make_survey(cls.owner, n=25)
        from .models_team import SurveyCollaborator
        SurveyCollaborator.objects.create(survey=cls.survey, user=cls.viewer, role="viewer")

    def setUp(self):
        cache.clear()

    def post(self, url, data, user=None):
        self.client.force_login(user or self.owner)
        return self.client.post(url, json.dumps(data), content_type="application/json")

    def test_status(self):
        self.client.force_login(self.owner)
        r = self.client.get("/kura/agent/status/", {"code": self.survey.code, "page": "data"})
        self.assertEqual(r.status_code, 200)
        j = r.json()
        self.assertTrue(j["can"]["export"])
        self.assertEqual(j["survey"]["code"], self.survey.code)
        self.client.force_login(self.outsider)
        r = self.client.get("/kura/agent/status/", {"code": self.survey.code})
        self.assertEqual(r.status_code, 404)

    def test_chat_builtin_when_ai_off(self):
        with mock.patch.object(agent_ai, "ai_enabled", return_value=False):
            r = self.post("/kura/agent/chat/", {"message": "pie chart of water source", "code": self.survey.code, "page": "data"})
        j = r.json()
        self.assertEqual(r.status_code, 200, j)
        self.assertEqual(j["model"], "kora-builtin")
        self.assertIn("built-in", j["notice"])
        self.assertEqual(j["blocks"][0]["type"], "chart")
        self.assertEqual(j["blocks"][0]["chart"], "pie")

    def test_chat_with_ai_plan(self):
        fake = {"say": "Here is your **breakdown**.", "mood": "proud", "suggestions": ["Split by sex"],
                "actions": [{"tool": "chart", "kind": "donut", "x": "region"},
                            {"tool": "table", "mode": "crosstab", "x": "water", "group_by": "sex"},
                            {"tool": "not_a_tool"}]}
        from hanns import ai as hai
        with mock.patch.object(hai, "is_enabled", return_value=True), \
                mock.patch.object(hai, "ai_json", return_value=(fake, "gemini-flash-latest")) as m:
            r = self.post("/kura/agent/chat/", {"message": "regions please", "code": self.survey.code, "page": "data",
                                                "history": [{"role": "user", "text": "hi"}]})
        j = r.json()
        self.assertEqual(j["model"], "gemini-flash-latest")
        self.assertEqual([b["type"] for b in j["blocks"]], ["chart", "table"])
        self.assertIsNone(j["notice"])
        sent_prompt = m.call_args[0][1]
        for name in SECRET_NAMES:
            self.assertNotIn(name, sent_prompt)
        self.assertLess(j["quotaLeft"], agent_ai.hourly_limit())

    def test_chat_falls_back_when_gemini_fails(self):
        from hanns import ai as hai
        with mock.patch.object(hai, "is_enabled", return_value=True), \
                mock.patch.object(hai, "ai_json", side_effect=hai.AIError("quota used up", status=429)):
            r = self.post("/kura/agent/chat/", {"message": "make a dashboard", "code": self.survey.code, "page": "data"})
        j = r.json()
        self.assertEqual(j["model"], "kora-builtin")
        self.assertIn("quota used up", j["notice"])
        self.assertEqual(j["blocks"][0]["type"], "dashboard")

    def test_chat_denied_for_outsider(self):
        r = self.post("/kura/agent/chat/", {"message": "hi", "code": self.survey.code}, user=self.outsider)
        self.assertEqual(r.status_code, 404)

    def test_exports(self):
        spec = {"tool": "chart", "kind": "bar", "x": "region"}
        for fmt, ctype in [("csv", "text/csv"), ("xlsx", "spreadsheetml"), ("html", "text/html")]:
            r = self.post("/kura/agent/export/", {"code": self.survey.code, "spec": spec, "format": fmt})
            self.assertEqual(r.status_code, 200, (fmt, r.content[:200]))
            self.assertIn(ctype, r["Content-Type"])
            self.assertIn(self.survey.code, r["X-Kora-Filename"])
        html = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "html",
                                                 "spec": {"tool": "dashboard", "auto": True}}).content.decode()
        self.assertIn("kora-export", html)
        self.assertIn("chart.umd.min.js", html)
        self.assertIn("leaflet", html)
        geo = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "geojson", "spec": {"tool": "map"}})
        self.assertEqual(json.loads(geo.content)["type"], "FeatureCollection")

    def test_row_level_export_has_every_row_and_no_names(self):
        r = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "csv", "spec": {"tool": "privacy"}})
        body = r.content.decode("utf-8-sig")
        self.assertEqual(body.count("\n"), 25 + 1 + 1)   # header + 25 + the duplicate
        for name in SECRET_NAMES:
            self.assertNotIn(name, body)
        x = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "xlsx", "spec": {"tool": "privacy"}})
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Safe data", "Privacy report", "Re-identification check"])

    def test_viewer_cannot_download_rows(self):
        r = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "csv", "spec": {"tool": "clean", "auto": True}},
                      user=self.viewer)
        self.assertEqual(r.status_code, 400)
        self.assertIn("permission", r.json()["error"])
        ok = self.post("/kura/agent/export/", {"code": self.survey.code, "format": "csv",
                                               "spec": {"tool": "chart", "x": "sex"}}, user=self.viewer)
        self.assertEqual(ok.status_code, 200)

    def test_kura_form_round_trip_shape(self):
        q = agent_ai.template_questionnaire("survey about a health clinic")
        r = self.post("/kura/agent/export/", {"format": "kuraform", "spec": {"tool": "questionnaire", **q}})
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("Questions", wb.sheetnames)
        rows = list(wb["Questions"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Type")
        self.assertIn("Choose one", [r_[0] for r_ in rows])
        show = [r_[6] for r_ in rows if r_[6]]
        self.assertIn("consent = yes", show)

    def test_apply_actions(self):
        q = agent_ai.template_questionnaire("event feedback for a workshop")
        r = self.post("/kura/agent/apply/", {"action": "create_survey", **q})
        j = r.json()
        self.assertTrue(j["ok"], j)
        new = Survey.objects.get(code=j["code"])
        self.assertEqual(new.owner, self.owner)
        self.assertEqual(new.draft_schema["questions"][1]["name"], "consent")

        r = self.post("/kura/agent/apply/", {"action": "add_questions", "code": self.survey.code,
                                             "questions": [{"type": "text", "label": "Extra", "name": "age"}]})
        self.assertTrue(r.json()["ok"])
        self.survey.refresh_from_db()
        names = [x["name"] for x in self.survey.draft_schema["questions"]]
        self.assertEqual(len(names), len(set(names)))

        dash = eng.run_tool(_ctx(self.owner, self.survey), {"tool": "dashboard", "auto": True})
        r = self.post("/kura/agent/apply/", {"action": "create_dashboard", "code": self.survey.code,
                                             "title": dash["title"], "theme": "forest", "tiles": dash["spec"]["tiles"]})
        j = r.json()
        self.assertTrue(j["ok"], j)
        from .models_dashboard import LiveDashboard
        board = LiveDashboard.objects.get(survey=self.survey)
        self.assertEqual(board.tiles.count(), len(dash["spec"]["tiles"]))
        self.assertTrue(j["url"].endswith(f"/dash/{board.id}/"))

        cl = eng.run_tool(_ctx(self.owner, self.survey), {"tool": "clean", "auto": True})
        r = self.post("/kura/agent/apply/", {"action": "save_pipeline", "code": self.survey.code,
                                             "name": "Kora test", "steps": cl["spec"]["steps"]})
        self.assertTrue(r.json()["ok"], r.json())
        pipe = self.survey.pipelines.get(name="Kora test")
        self.assertEqual(pipe.steps.count(), len(cl["spec"]["steps"]))
        # and it runs in the real studio executor
        from .models import CleaningRun
        from .pipeline_engine import PipelineExecutor
        run = CleaningRun.objects.create(pipeline=pipe)
        PipelineExecutor(run).execute()
        run.refresh_from_db()
        self.assertEqual(run.status, "complete")

        r = self.post("/kura/agent/apply/", {"action": "save_dataset", "code": self.survey.code, "spec": {}})
        self.assertTrue(r.json()["ok"], r.json())
        ds = self.survey.datasets.latest("id")
        self.assertNotIn("full_name", ds.columns)

        r = self.post("/kura/agent/apply/", {"action": "set_state", "code": self.survey.code, "state": "paused"})
        self.assertTrue(r.json()["ok"])
        self.survey.refresh_from_db()
        self.assertEqual(self.survey.state, "paused")

    def test_viewer_cannot_apply(self):
        for action in ("add_questions", "save_pipeline", "save_dataset", "set_state", "create_dashboard"):
            r = self.post("/kura/agent/apply/", {"action": action, "code": self.survey.code,
                                                 "questions": [{"type": "text", "label": "x"}],
                                                 "steps": [{"operation": "trim_text", "config": {"field": "village"}}],
                                                 "state": "closed", "tiles": [{"kind": "kpi", "spec": {}}]},
                          user=self.viewer)
            self.assertEqual(r.status_code, 403, action)


def _ctx(user, survey, page="data"):
    from .agent_views import build_context
    return build_context(user, survey.code, page)


@override_settings(STATICFILES_STORAGE="django.contrib.staticfiles.storage.StaticFilesStorage")
class AgentMiddlewareTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("mw", password="pw")
        cls.outsider = User.objects.create_user("mw2", password="pw")
        cls.survey = make_survey(cls.owner, n=3)

    def test_injected_on_owner_pages(self):
        self.client.force_login(self.owner)
        for url in (f"/kura/{self.survey.code}/build/", f"/kura/{self.survey.code}/data/", "/kura/"):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
            body = r.content.decode()
            self.assertIn('id="kora-config"', body, url)
            self.assertIn("kura/agent/kora.js", body)
            self.assertLess(body.index("kora-config"), body.rindex("</body>"))

    def test_config_is_safe_json(self):
        self.survey.title = '</script><script>alert(1)</script>'
        self.survey.save()
        self.client.force_login(self.owner)
        body = self.client.get(f"/kura/{self.survey.code}/data/").content.decode()
        start = body.index('id="kora-config"')
        chunk = body[start:start + 2000]
        self.assertNotIn("<script>alert(1)", chunk)

    def test_not_on_public_runner_or_for_anonymous(self):
        r = self.client.get(f"/kura/{self.survey.code}/")
        self.assertNotIn(b"kora-config", r.content)
        self.client.force_login(self.owner)
        r = self.client.get(f"/kura/{self.survey.code}/")
        self.assertNotIn(b"kora-config", r.content)

    def test_opt_out_and_switch(self):
        self.client.force_login(self.owner)
        r = self.client.get(f"/kura/{self.survey.code}/data/?kora=0")
        self.assertNotIn(b"kora-config", r.content)
        with override_settings(KURA_AGENT_ENABLED=False):
            r = self.client.get(f"/kura/{self.survey.code}/data/")
            self.assertNotIn(b"kora-config", r.content)
