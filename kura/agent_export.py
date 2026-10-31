"""
kura/agent_export.py — turn anything Kora made into a file.

Every export RE-RUNS the tool on the server from its normalised spec, with
the caller's permissions, instead of trusting numbers sent by the browser.
That is what lets a chat preview show 60 rows while the download carries
all 12,000, and it means a tampered request can't smuggle in data the user
isn't allowed to see.

Formats
  csv      UTF-8 with BOM so Excel opens accents and Wolof/Mandinka text correctly
  xlsx     styled workbook (one sheet per table; privacy report sheet included)
  html     standalone page: Chart.js + Leaflet from CDNs, data embedded
  geojson  map points
  json     questionnaire fragment
  kuraform the builder's own Excel format — re-imports with ⇪ Import
"""

from __future__ import annotations

import io
import json
import re

import pandas as pd
from django.template.loader import render_to_string
from django.utils import timezone

from . import agent_engine as eng

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "kuraform": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "html": "text/html; charset=utf-8",
    "geojson": "application/geo+json",
    "json": "application/json",
}
ROW_LEVEL = {"clean", "privacy"}


class ExportError(Exception):
    pass


def _slug(text, fallback="kora"):
    s = re.sub(r"[^A-Za-z0-9]+", "-", str(text or "")).strip("-").lower()[:50]
    return s or fallback


def _sheet_name(text, used):
    base = re.sub(r"[\[\]\:\*\?\/\\]", " ", str(text or "Sheet")).strip()[:28] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:25]} {n}"
        n += 1
    used.add(name.lower())
    return name


# ─────────────────────────── tables out of blocks ───────────────────────────

def _chart_frame(block):
    labels = block.get("labels") or []
    series = block.get("series") or []
    if block.get("chart") == "scatter":
        rows = []
        for s in series:
            for p in s.get("data") or []:
                rows.append({"Series": s.get("name"), block.get("x_label") or "X": p.get("x"),
                             block.get("y_label") or "Y": p.get("y")})
        return pd.DataFrame(rows)
    data = {block.get("x_label") or "Category": labels}
    for s in series:
        data[str(s.get("name") or "Value")] = list(s.get("data") or [])[: len(labels)]
    if block.get("percent") and len(series) == 1:
        data["Percent"] = block["percent"][: len(labels)]
    return pd.DataFrame(data)


def block_frames(block) -> list[tuple[str, pd.DataFrame]]:
    t = block.get("type")
    title = block.get("title") or t
    if t == "chart":
        return [(title, _chart_frame(block))]
    if t == "table":
        return [(title, pd.DataFrame(block.get("rows") or [], columns=block.get("columns") or None))]
    if t == "cards":
        return [(title or "Key numbers", pd.DataFrame([{"Measure": c.get("label"), "Value": c.get("value"),
                                                        "Note": c.get("hint")} for c in block.get("cards") or []]))]
    if t == "shape":
        return [(title, pd.DataFrame([{"Label": i.get("label"), "Value": i.get("value"), "Percent": i.get("pct")}
                                      for i in block.get("items") or []]))]
    if t == "map":
        return [(title, pd.DataFrame([{"Latitude": p.get("lat"), "Longitude": p.get("lng"),
                                       "Group": p.get("group", "")} for p in block.get("points") or []]))]
    if t == "profile":
        return [("Columns", pd.DataFrame([{"Column": r["name"], "Question": r["label"], "Type": r["type"],
                                           "Answered": r["answered"], "Missing %": r["missing_pct"],
                                           "Summary": r["summary"], "Personal data": r["pii"]}
                                          for r in block.get("rows") or []])),
                ("Issues", pd.DataFrame([{"Level": i["level"], "Issue": i["text"]} for i in block.get("issues") or []]))]
    if t == "overview":
        return [("Overview", pd.DataFrame([{"Measure": s["label"], "Value": s["value"]} for s in block.get("stats") or []])),
                ("Last 14 days", pd.DataFrame([{"Day": d["day"], "Responses": d["n"]} for d in block.get("trend") or []]))]
    if t == "dashboard":
        out = []
        for it in block.get("items") or []:
            out.extend(block_frames(it["block"]))
        return out
    if t == "lint":
        return [("Questionnaire check", pd.DataFrame([{"Level": i["level"], "Question": i.get("field", ""),
                                                        "Issue": i["text"]} for i in block.get("issues") or []]))]
    return []


def full_frames(ctx, spec) -> tuple[list[tuple[str, pd.DataFrame]], dict]:
    """Row-level frames with every row (not the chat preview)."""
    tool = spec.get("tool")
    if tool == "clean":
        steps = eng.clean_steps(ctx, spec)
        df, report = eng.run_clean(ctx, steps)
        df = df[[c for c in eng.public_columns(df) if c != "_uuid"]]
        steps_df = pd.DataFrame([{"Step": r["name"], "Operation": r["operation"], "Status": r["status"],
                                  "Rows before": r.get("rows_before"), "Rows after": r.get("rows_after"),
                                  "Values changed": r.get("changes"), "Error": r.get("error", "")} for r in report])
        return [("Cleaned data", _flatten(df)), ("Cleaning steps", steps_df)], {"title": spec.get("title") or "cleaned-data"}
    if tool == "privacy":
        df, plan, kinfo = eng.run_privacy(ctx, spec)
        rep = pd.DataFrame([{"Column": p["column"], "Question": p["label"], "Detected": p["level"],
                             "Why": p["reason"], "What Kora did": eng.ACTION_TEXT.get(p["action"], p["action"])}
                            for p in plan if p["column"] != "_submission_id"])
        k = pd.DataFrame([{"Check": "k-anonymity target (k)", "Value": kinfo["k"]},
                          {"Check": "Quasi-identifiers kept", "Value": ", ".join(kinfo["quasi"]) or "none"},
                          {"Check": "Smallest group", "Value": kinfo["min_group"]},
                          {"Check": "Rows in groups smaller than k", "Value": kinfo["at_risk"]},
                          {"Check": "Rows removed for safety", "Value": kinfo["suppressed"]}])
        return [("Safe data", _flatten(df)), ("Privacy report", rep), ("Re-identification check", k)], \
            {"title": spec.get("title") or "privacy-safe-data"}
    if tool == "table":
        df, meta = eng.table_frame(ctx, spec, preview=False)
        return [(spec.get("title") or "Table", _flatten(df))], {"title": spec.get("title") or "table"}
    return [], {}


def _flatten(df):
    out = df.copy()
    for c in out.columns:
        if out[c].apply(lambda v: isinstance(v, (list, dict))).any():
            out[c] = out[c].apply(lambda v: ("|".join(map(str, v)) if isinstance(v, list)
                                             else json.dumps(v, ensure_ascii=False)) if isinstance(v, (list, dict)) else v)
    return out


# ─────────────────────────── writers ───────────────────────────

def write_csv(frames) -> bytes:
    if not frames:
        raise ExportError("Nothing to export.")
    df = frames[0][1]
    return df.to_csv(index=False).encode("utf-8-sig")


def write_xlsx(frames, title="Kora export") -> bytes:
    if not frames:
        raise ExportError("Nothing to export.")
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    used = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in frames:
            sheet = _sheet_name(name, used)
            safe_df = df.copy()
            for c in safe_df.columns:
                if safe_df[c].dtype == object:
                    safe_df[c] = safe_df[c].apply(
                        lambda v: v if not isinstance(v, str) or not v[:1] in ("=", "+", "-", "@") or _looks_number(v)
                        else "'" + v)
            safe_df.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            head_fill = PatternFill("solid", fgColor="2B2F77")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for i, col in enumerate(safe_df.columns, start=1):
                sample = [str(col)] + [str(v) for v in safe_df[col].head(200).tolist()]
                width = min(60, max(8, max(len(s) for s in sample) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            ws.row_dimensions[1].height = 30
        xw.book.properties.title = title[:250]
        xw.book.properties.creator = "Kora · Kura"
    return buf.getvalue()


def _looks_number(v):
    try:
        float(v)
        return True
    except ValueError:
        return False


def write_geojson(block) -> bytes:
    feats = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [p["lng"], p["lat"]]},
              "properties": {"group": p.get("group", ""), "row": p.get("row")}} for p in block.get("points") or []]
    return json.dumps({"type": "FeatureCollection", "features": feats}, ensure_ascii=False).encode("utf-8")


def write_html(blocks_or_board, title, survey_title="", dark=True) -> bytes:
    if isinstance(blocks_or_board, dict) and blocks_or_board.get("type") == "dashboard":
        board = blocks_or_board
        items = board.get("items") or []
        palette = board.get("palette")
        tokens = board.get("theme_tokens") or {}
        dark = board.get("dark", True)
    else:
        items = [{"x": 0, "y": i * 4, "w": 12, "h": 4, "block": b} for i, b in enumerate(blocks_or_board)]
        palette, tokens = None, {}
    payload = {"title": title, "survey": survey_title, "items": items,
               "palette": palette or ["#22d3ee", "#ff3f98", "#ff8a1f", "#8b5cf6", "#22c55e", "#e8482b", "#38bdf8", "#f472b6"],
               "tokens": tokens, "dark": bool(dark),
               "generated": timezone.localtime().strftime("%d %b %Y, %H:%M")}
    needs_map = any(it["block"].get("type") == "map" for it in items)
    html = render_to_string("kura/agent_export.html", {
        "title": title, "payload": payload, "needs_map": needs_map,
    })
    return html.encode("utf-8")


# ─────────────────────────── Kura Form workbook ───────────────────────────
# Mirrors buildWorkbook()/_questionRow() in templates/kura/builder.html so
# the file re-imports with the builder's ⇪ Import, rows and logic intact.

FRIENDLY_TYPES = {
    "section": "Section header", "note": "Note / disclaimer", "text": "Short text", "long_text": "Paragraph",
    "integer": "Whole number", "decimal": "Decimal number", "select_one": "Choose one",
    "select_multiple": "Choose many", "rank": "Rank options", "likert": "Likert scale", "rating": "Rating",
    "date": "Date", "time": "Time", "datetime": "Date & time", "gps": "GPS location",
    "mapper": "Map (track / area / draw)", "photo": "Photo", "barcode": "Barcode / QR",
    "calculate": "Calculated (hidden)", "repeat": "Repeat group",
}
CHOICE_Q = {"select_one", "select_multiple", "rank", "likert"}
CMP_TO_TEXT = {"eq": "=", "ne": "≠", "gt": ">", "gte": ">=", "lt": "<", "lte": "<=", "contains": "contains",
               "not_contains": "doesn't contain", "selected": "has selected", "not_selected": "hasn't selected",
               "answered": "is answered", "not_answered": "is blank", "between": "between", "matches": "matches"}
KF_HEADERS = ["Type", "Question", "Name (data column)", "Required", "Choices", "Hint / help", "Show when",
              "Filter choices by", "Min", "Max", "Pattern (regex)", "Must be true (rule)", "Warn when false",
              "Warning message", "Error message", "Calculation", "Indicator"]


def _show_when(rel):
    if not isinstance(rel, dict) or not rel.get("rules"):
        return ""
    if any(isinstance(r, dict) and isinstance(r.get("rules"), list) for r in rel["rules"]):
        return "(complex logic — edit in the builder)"
    parts = []
    for r in rel["rules"]:
        if r.get("cmp") == "answered":
            parts.append(f"{r['q']} is answered")
        elif r.get("cmp") == "not_answered":
            parts.append(f"{r['q']} is blank")
        else:
            v = ", ".join(map(str, r["value"])) if isinstance(r.get("value"), list) else ("" if r.get("value") is None else r["value"])
            parts.append(f"{r['q']} {CMP_TO_TEXT.get(r.get('cmp'), '=')} {v}")
    return (" or " if rel.get("op") == "or" else " and ").join(parts)


def _choices_cell(q):
    cs = q.get("choices") or []
    if not cs:
        return ""
    if any(c.get("parent") not in (None, "") for c in cs):
        return "(see the Choices sheet)"
    return "\n".join(f"{c['value']} = {c.get('label') or c['value']}" for c in cs)


def _kf_row(q):
    v = q.get("validate") or {}
    rep = q.get("repeat") or {}
    return [
        FRIENDLY_TYPES.get(q.get("type"), q.get("type")), q.get("label") or "", q.get("name") or "",
        "Yes" if q.get("required") else "", _choices_cell(q) if q.get("type") in CHOICE_Q else "",
        q.get("hint") or "", _show_when(q.get("relevant")), (q.get("cascade") or {}).get("parent", ""),
        rep.get("min", "") if q.get("type") == "repeat" else v.get("min", ""),
        rep.get("max", "") if q.get("type") == "repeat" else v.get("max", ""),
        v.get("regex", ""), v.get("expr", ""), v.get("warn_expr", ""), v.get("warn_message", ""),
        v.get("message", ""), rep.get("count_expr", "") if q.get("type") == "repeat" else q.get("calc", ""), "",
    ]


def write_kura_form(title, description, questions) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "📖 Read me"
    for line in [
        "KURA FORM — fill this workbook, then Import it in the form builder.", "",
        "Drafted by Kora, Kura's AI assistant. Edit freely: one row per question in the Questions sheet.",
        "• Type: Choose one, Choose many, Short text, Paragraph, Whole number, Decimal number, Rating, Likert scale,",
        "  Rank options, Date, Time, Date & time, GPS location, Photo, Barcode / QR, Note / disclaimer,",
        "  Section header, Calculated (hidden), Repeat group, End repeat.",
        "• Choices: one per line as  value = Label.",
        "• Show when: e.g.  consent = yes  ·  age >= 18  ·  symptoms has selected fever.",
        "• Repeat group rows continue until End repeat.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 100

    wf = wb.create_sheet("Form")
    wf.append(["Title", title or ""])
    wf.append(["Description", description or ""])
    wf.column_dimensions["A"].width = 14
    wf.column_dimensions["B"].width = 70

    wq = wb.create_sheet("Questions")
    wq.append(KF_HEADERS)
    choice_rows = []
    for q in questions or []:
        wq.append(_kf_row(q))
        if q.get("type") in CHOICE_Q and _choices_cell(q) == "(see the Choices sheet)":
            for c in q.get("choices") or []:
                choice_rows.append([q["name"], c["value"], c.get("label", ""), c.get("parent", ""), ""])
        if q.get("type") == "repeat":
            for c in q.get("children") or []:
                wq.append(_kf_row(c))
            wq.append(["End repeat"] + [""] * 16)
    fill = PatternFill("solid", fgColor="2B2F77")
    for cell in wq[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    widths = [18, 46, 18, 9, 30, 32, 28, 16, 7, 7, 18, 22, 20, 24, 24, 18, 12]
    for i, w in enumerate(widths):
        wq.column_dimensions[chr(65 + i)].width = w
    for row in wq.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wq.freeze_panes = "A2"
    if choice_rows:
        wc = wb.create_sheet("Choices")
        wc.append(["Question name", "Value", "Label", "Parent value", "Zones"])
        for r in choice_rows:
            wc.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────── entry point ───────────────────────────

def export(ctx, spec: dict, fmt: str):
    """Returns (bytes, filename, content_type). Raises ExportError / eng.ToolError."""
    tool = (spec or {}).get("tool")
    if tool not in eng.TOOLS:
        raise ExportError("Unknown block.")
    if fmt not in CONTENT_TYPES:
        raise ExportError("Unknown file type.")
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    code = ctx.survey.code if ctx.survey else "kura"
    survey_title = ctx.survey.title if ctx.survey else ""

    if tool == "questionnaire":
        qs = eng.sanitize_questions(spec.get("questions"))
        title = spec.get("title") or "Draft questionnaire"
        base = f"{_slug(title)}-{stamp}"
        if fmt == "json":
            data = json.dumps({"title": title, "description": spec.get("description") or "",
                               "schema": {"settings": {"anonymous": True, "languages": ["en"]}, "questions": qs}},
                              ensure_ascii=False, indent=2).encode("utf-8")
            return data, base + ".json", CONTENT_TYPES["json"]
        if fmt in ("kuraform", "xlsx"):
            return write_kura_form(title, spec.get("description") or "", qs), base + "_kura_form.xlsx", CONTENT_TYPES["xlsx"]
        raise ExportError("Questionnaires download as Kura Form (Excel) or JSON.")

    if tool in ROW_LEVEL or (tool == "table" and spec.get("mode") == "rows"):
        if not ctx.can.get("export"):
            raise ExportError("You don't have permission to download row-level data for this survey.")

    if tool in ROW_LEVEL or tool == "table":
        frames, meta = full_frames(ctx, spec)
        title = meta.get("title") or tool
        base = f"{code}-{_slug(title)}-{stamp}"
        if fmt == "csv":
            return write_csv(frames), base + ".csv", CONTENT_TYPES["csv"]
        if fmt == "xlsx":
            return write_xlsx(frames, title), base + ".xlsx", CONTENT_TYPES["xlsx"]
        if fmt == "html" and tool == "table":
            block = eng.run_tool(ctx, spec)
            df = frames[0][1]
            cols, rows = eng.frame_records(df, 5000)
            block = {**block, "columns": cols, "rows": rows}
            return write_html([block], block.get("title") or title, survey_title), base + ".html", CONTENT_TYPES["html"]
        raise ExportError("This one downloads as CSV or Excel.")

    block = eng.run_tool(ctx, spec)
    if block.get("type") == "error":
        raise ExportError(block.get("text") or "Could not rebuild this block.")
    title = block.get("title") or tool
    base = f"{code}-{_slug(title)}-{stamp}"
    if fmt == "html":
        payload = block if block.get("type") == "dashboard" else [block]
        return write_html(payload, title, survey_title), base + ".html", CONTENT_TYPES["html"]
    if fmt == "geojson":
        if block.get("type") != "map":
            raise ExportError("Only maps download as GeoJSON.")
        return write_geojson(block), base + ".geojson", CONTENT_TYPES["geojson"]
    frames = block_frames(block)
    if fmt == "csv":
        return write_csv(frames), base + ".csv", CONTENT_TYPES["csv"]
    if fmt == "xlsx":
        return write_xlsx(frames, title), base + ".xlsx", CONTENT_TYPES["xlsx"]
    raise ExportError("Unsupported file type for this block.")
